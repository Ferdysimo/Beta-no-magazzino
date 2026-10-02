import asyncio
import sys
import threading
from io import BytesIO
from pathlib import Path

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook


BACKEND_DIR = Path(__file__).resolve().parents[1]
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))

from app.routers import analysis as analysis_router
from app.services import analysis as analysis_service


def test_analysis_workbook_builder_returns_seeked_xlsx(monkeypatch):
    locale_calls = []

    def fake_locale_writer(workbook, restaurant_data, data, used_titles):
        locale_calls.append(restaurant_data["id"])
        workbook.create_sheet("Locale")

    def fake_totals_writer(workbook, restaurants, selected_year):
        assert selected_year == 2026
        assert [restaurant["id"] for restaurant in restaurants] == ["flaminio"]
        workbook.create_sheet("Totali")

    monkeypatch.setattr(analysis_service, "_write_analysis_locale_sheet", fake_locale_writer)
    monkeypatch.setattr(analysis_service, "_write_totali_sheet_for_analysis", fake_totals_writer)

    output = analysis_service._build_analysis_workbook_bytes(
        {"restaurants": [{"id": "flaminio"}]},
        2026,
    )

    assert locale_calls == ["flaminio"]
    assert output.tell() == 0
    assert load_workbook(output, read_only=True).sheetnames == ["Locale", "Totali"]


def test_analysis_workbook_runs_off_event_loop_and_rejects_parallel_export(monkeypatch):
    started = threading.Event()
    release = threading.Event()
    worker_threads = []
    data = {
        "restaurants": [],
        "integrity": {"errors": [], "warnings": [], "warning_counts": {}},
    }

    async def fake_build_data(selected_year):
        assert selected_year == 2026
        return data

    def slow_workbook_builder(received_data, selected_year):
        worker_threads.append(threading.get_ident())
        assert received_data is data
        assert selected_year == 2026
        started.set()
        assert release.wait(timeout=2), "test did not release workbook worker"
        return BytesIO(b"xlsx")

    monkeypatch.setattr(analysis_router, "_build_annual_analysis_data", fake_build_data)
    monkeypatch.setattr(analysis_router, "_build_analysis_workbook_bytes", slow_workbook_builder)

    async def exercise():
        main_thread = threading.get_ident()
        first_export = asyncio.create_task(
            analysis_router.export_analisi_mensile_excel(
                year=2026,
                token_data={"role": "admin", "username": "Admin"},
            )
        )

        for _ in range(100):
            if started.is_set():
                break
            await asyncio.sleep(0.005)
        assert started.is_set(), "workbook generation did not start"

        # If OpenPyXL were still running on the event-loop thread, execution
        # could not reach this assertion until the blocking builder returned.
        await asyncio.sleep(0)
        assert not first_export.done()
        assert len(worker_threads) == 1
        assert worker_threads[0] != main_thread

        with pytest.raises(HTTPException) as exc_info:
            await analysis_router.export_analisi_mensile_excel(
                year=2026,
                token_data={"role": "admin", "username": "Admin"},
            )
        assert exc_info.value.status_code == 429
        assert "già in generazione" in exc_info.value.detail

        release.set()
        response = await asyncio.wait_for(first_export, timeout=2)
        assert response.headers["content-disposition"] == (
            'attachment; filename="analisi_mensile_2026.xlsx"'
        )

        # The slot must be released after success so a later export can start.
        release.clear()
        release.set()
        second_response = await analysis_router.export_analisi_mensile_excel(
            year=2026,
            token_data={"role": "admin", "username": "Admin"},
        )
        assert second_response.status_code == 200

    try:
        asyncio.run(exercise())
    finally:
        release.set()
        analysis_router._release_analysis_export()


def test_analysis_export_slot_is_released_after_failure(monkeypatch):
    data = {
        "restaurants": [],
        "integrity": {"errors": [], "warnings": [], "warning_counts": {}},
    }

    async def fake_build_data(selected_year):
        return data

    def broken_workbook_builder(received_data, selected_year):
        raise RuntimeError("synthetic workbook failure")

    monkeypatch.setattr(analysis_router, "_build_annual_analysis_data", fake_build_data)
    monkeypatch.setattr(analysis_router, "_build_analysis_workbook_bytes", broken_workbook_builder)

    async def exercise():
        with pytest.raises(RuntimeError, match="synthetic workbook failure"):
            await analysis_router.export_analisi_mensile_excel(
                year=2026,
                token_data={"role": "admin", "username": "Admin"},
            )
        assert analysis_router._reserve_analysis_export() is True
        analysis_router._release_analysis_export()

    try:
        asyncio.run(exercise())
    finally:
        analysis_router._release_analysis_export()
